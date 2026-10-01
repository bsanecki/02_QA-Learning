from openpyxl import load_workbook
import random
import time
import os


# ========================================
# USTAWIENIA ŚCIEŻEK
# ========================================

# WSTAW TUTAJ ŚCIEŻKĘ DO PLIKU EXCEL
plik = 'WSTAW_TUTAJ_SCIEZKE/dane_biblioteki.xlsx'

# WSTAW TUTAJ ŚCIEŻKĘ DO PLIKU Z KSIĄŻKAMI
SCIEZKA_ZBIORU_KSIAZEK = 'WSTAW_TUTAJ_SCIEZKE/zbior_ksiazek.txt'

# WSTAW TUTAJ ŚCIEŻKĘ DO PLIKU Z WYKORZYSTANYMI ID
SCIEZKA_ZBIORU_ID = 'WSTAW_TUTAJ_SCIEZKE/wykorzystane_id.txt'

# WSTAW TUTAJ ŚCIEŻKĘ DO PLIKU Z WYPOŻYCZONYMI KSIĄŻKAMI
SCIEZKA_WYPOZYCZONYCH_KSIAZEK = 'WSTAW_TUTAJ_SCIEZKE/ksiazki_wypozyczone.txt'

# WSTAW TUTAJ ŚCIEŻKĘ DO FOLDERU NA WYPOŻYCZONE KSIĄŻKI
FOLDER_WYPOZYCZONYCH_KSIAZEK = 'WSTAW_TUTAJ_SCIEZKE/Wypozyczone_ksiazki'


# ========================================
# WCZYTANIE PLIKU EXCEL
# ========================================

wb = load_workbook(plik)
ws = wb.active


class panel_dzialania_admina:

    def __init__(self):
        self.sciezka_zbioru_ksiazek = SCIEZKA_ZBIORU_KSIAZEK
        self.sciezka_do_zbioru_id = SCIEZKA_ZBIORU_ID
        self.sciezka_do_wypozyczonych_ksiazek = SCIEZKA_WYPOZYCZONYCH_KSIAZEK
        self.folder_wypo_ks_do_id = FOLDER_WYPOZYCZONYCH_KSIAZEK

    def menu_programu(self):
        print("""
=== SYSTEM BIBLIOTECZNY ===
1. Dodaj książkę
2. Usuń książkę
3. Szukaj książki
4. Wyświetl wszystkie książki
5. Dodaj użytkownika
6. Wypożycz książkę
7. Zwróć książkę
0. Wyjdź
===========================
Wybierz numer operacji (0-7):
""")

    def dodaj_ksiazke(self):
        print("Podaj tytuł książki")
        tytul = input()
        self.tytul = tytul

        with open(self.sciezka_zbioru_ksiazek, 'r', encoding='utf-8') as plik1:
            zawartosc = plik1.read()

        if self.tytul in zawartosc:
            print("Książka znajduje sie w zbiorze")
            return
        else:
            with open(self.sciezka_zbioru_ksiazek, 'a', encoding='utf-8') as plik:
                plik.write("\n" + self.tytul)
                print("Książka została dodana do zbioru!")
                return

    def usun_ksiazke(self):
        print("Podaj tytuł książki, którą chcesz usunąć z zbioru!")
        wybor_ksiazki = input().strip()

        with open(self.sciezka_zbioru_ksiazek, 'r', encoding='utf-8') as plik:
            linie = plik.readlines()

        linie_po_usunieciu = [
            linia for linia in linie
            if linia.strip() != wybor_ksiazki
        ]

        if len(linie) == len(linie_po_usunieciu):
            print("Nie ma takiej książki w zbiorze!")
            return
        else:
            with open(self.sciezka_zbioru_ksiazek, 'w', encoding='utf-8') as plik:
                plik.writelines(linie_po_usunieciu)

            print(f"Książka '{wybor_ksiazki}' została usunięta.")
            return

    def szukaj_ksiazki(self):
        print("Podaj tytuł książki którą szukasz!")
        wybor = input()

        with open(self.sciezka_zbioru_ksiazek, 'r', encoding='utf-8') as plik:
            linie = [linia.strip() for linia in plik]

        if wybor in linie:
            print("Książka znajduje się w zbiorze biblioteki!")
        else:
            print("W zbiorze nie ma tej książki!")

    def wyswietl_wszystkie_ksiazki(self):

        with open(self.sciezka_zbioru_ksiazek, 'r', encoding='utf-8') as plik:
            ksiazki = [
                linia.strip()
                for linia in plik
                if linia.strip()
            ]

        if not ksiazki:
            print("Brak książek w zbiorze.")
        else:
            print("Oto zbiór książek w bibliotece!")

            for ksiazka in ksiazki:
                print(ksiazka)

    def dodaj_uzytkownika(self):
        print("Podaj imię i nazwisko")
        imie_nazwisko = input()

        print("Podaj date urodzenia np. 01.01.1999")
        data = input()

        print("Adres zamieszkania")
        adres = input()

        wiersz = 1

        while ws[f'A{wiersz}'].value is not None:
            wiersz += 1

        ws[f'A{wiersz}'] = imie_nazwisko
        ws[f'B{wiersz}'] = data
        ws[f'C{wiersz}'] = adres

        with open(
            self.sciezka_do_zbioru_id,
            'r+',
            encoding='utf-8'
        ) as plik3:

            zawartosc = plik3.read()

        id = None

        while True:
            id = random.randint(500, 100000)

            if str(id) not in zawartosc:
                break

        plik3.close()

        with open(
            self.sciezka_do_zbioru_id,
            'a',
            encoding='utf-8'
        ) as plik4:

            plik4.write(str(id) + '\n')

            plik4.close()

        ws[f'D{wiersz}'] = id

        wb.save(plik)

        print(
            imie_nazwisko,
            "został zapisany do systemu z numerem id =",
            id,
            "."
        )

    def wypozycz_ksiazke(self):

        with open(
            self.sciezka_do_zbioru_id,
            'r',
            encoding='utf-8'
        ) as zbior_id:

            wszystkie_id = [
                linia.strip()
                for linia in zbior_id
                if linia.strip()
            ]

        print("Podaj id użytkownika:")
        id = input().strip()

        if id not in wszystkie_id:
            print("Użytkownik o podanym ID nie istnieje")
            return

        print(
            "Podaj dokładny tytuł książki, "
            "którą użytkownik chce wypożyczyć:"
        )

        wybor_ksiazki = input().strip()

        with open(
            self.sciezka_do_wypozyczonych_ksiazek,
            'r',
            encoding='utf-8'
        ) as zbior_wp_ks:

            wypozyczone = [
                linia.strip()
                for linia in zbior_wp_ks
                if linia.strip()
            ]

        if wybor_ksiazki in wypozyczone:
            print("Książka jest już wypożyczona!")
            return

        with open(
            self.sciezka_zbioru_ksiazek,
            'r',
            encoding='utf-8'
        ) as zbior_ksiazek:

            wszystkie_ksiazki = [
                linia.strip()
                for linia in zbior_ksiazek
                if linia.strip()
            ]

        if wybor_ksiazki not in wszystkie_ksiazki:
            print(
                "Podana książka nie znajduje się "
                "w zbiorze biblioteki!"
            )
            return

        os.makedirs(
            self.folder_wypo_ks_do_id,
            exist_ok=True
        )

        sciezka_pliku = os.path.join(
            self.folder_wypo_ks_do_id,
            f"{id}.txt"
        )

        if os.path.isfile(sciezka_pliku):

            with open(
                sciezka_pliku,
                'a',
                encoding='utf-8'
            ) as zbior_ks_uzutkownika:

                zbior_ks_uzutkownika.write(
                    wybor_ksiazki + '\n'
                )

        else:

            with open(
                sciezka_pliku,
                'w',
                encoding='utf-8'
            ) as nowy_plik:

                nowy_plik.write(
                    wybor_ksiazki + '\n'
                )

        with open(
            self.sciezka_do_wypozyczonych_ksiazek,
            'a',
            encoding='utf-8'
        ) as plik_wyp:

            plik_wyp.write(
                wybor_ksiazki + '\n'
            )

        print(
            f"Książka '{wybor_ksiazki}' "
            f"została wypożyczona użytkownikowi: {id}"
        )

    def zwroc_ksiazke(self):

        print("Podaj id użytkownika")
        id = input().strip()

        sciezka_pliku = os.path.join(
            self.folder_wypo_ks_do_id,
            f"{id}.txt"
        )

        if not os.path.isfile(sciezka_pliku):
            print("Użytkownik nie ma wypożyczonych książek.")
            return

        print("Podaj tytuł książki do oddania:")
        tytul_ks = input().strip()

        with open(
            sciezka_pliku,
            'r',
            encoding='utf-8'
        ) as plik:

            ksiazki = [
                linia.strip()
                for linia in plik
            ]

        if tytul_ks not in ksiazki:
            print("Użytkownik nie wypożyczył tej książki.")
            return

        ksiazki.remove(tytul_ks)

        with open(
            sciezka_pliku,
            'w',
            encoding='utf-8'
        ) as plik:

            for k in ksiazki:
                plik.write(k + '\n')

        print(
            f"Książka '{tytul_ks}' "
            f"została zwrócona przez użytkownika {id}."
        )

    def wyjdz(self):
        print("Koniec programu.")


panel = panel_dzialania_admina()


while True:

    time.sleep(1)

    panel.menu_programu()

    operacja = input().strip()

    if operacja == '0':
        panel.wyjdz()
        break

    elif operacja == '1':
        panel.dodaj_ksiazke()

    elif operacja == '2':
        panel.usun_ksiazke()

    elif operacja == '3':
        panel.szukaj_ksiazki()

    elif operacja == '4':
        panel.wyswietl_wszystkie_ksiazki()

    elif operacja == '5':
        panel.dodaj_uzytkownika()

    elif operacja == '6':
        panel.wypozycz_ksiazke()

    elif operacja == '7':
        panel.zwroc_ksiazke()

    else:
        print(
            "Nieprawidłowy wybór! "
            "Wprowadź numer od 0 do 7."
        )

        time.sleep(1)
